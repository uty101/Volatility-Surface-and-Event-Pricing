"""Long history pull (PLAN 1.1, instruction 01 step 1.1).

Each stage is independent and rerunnable, writes its files under data/raw/ and rewrites its entries in
data/raw/MANIFEST.json. Network access happens only here: stages take a session object (anything with
get(url) returning a response with .content and .raise_for_status()) and, for prices, a ticker factory,
so tests can drive the parsers with canned bytes.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import platform
import sys
from datetime import UTC, date, datetime, timedelta
from importlib.metadata import version
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from volsurf.config import load_config
from volsurf.snapshots import NY

RAW = ROOT / "data" / "raw"
LIBS = ["yfinance", "pandas", "numpy", "pyarrow", "openpyxl", "requests", "curl_cffi"]
CSV = {"index": False, "float_format": "%.10g", "lineterminator": "\n", "encoding": "utf-8"}

CBOE_CSV_URL = "https://cdn-api.cboe.com/api/global/us_indices/daily_prices/{name}_History.csv"
CBOE_JSON_URL = "https://cdn-api.cboe.com/api/global/delayed_quotes/charts/historical/_{name}.json"
CBOE_CSV = ["VIX", "VIX9D", "VIX3M", "VIX6M"]
CBOE_JSON = ["PUT", "BXM", "COR3M"]
FRED_URL = "https://fred.stlouisfed.org/graph/fredgraph.csv?id={id}"
FRED_IDS = ["DGS1MO", "DGS3MO", "DGS6MO", "DGS1", "DTB3"]
SSGA_URL = ("https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/"
            "holdings-daily-us-en-xlk.xlsx")
PRICES_START = "1989-12-01"
INDEX_TICKERS = ["^GSPC", "^SP500TR"]
XLK_WEIGHT_GATE = (0.97, 1.01)  # instruction 01, step 1.1 ssga gate


class GateError(RuntimeError):
    """A gate in the instruction failed its threshold (rule 4)."""


# ---------------------------------------------------------------- parsers (pure, no network)

def cboe_json_to_frame(raw: bytes) -> pd.DataFrame:
    """Cboe chart JSON -> date, close with ISO dates, sorted."""
    data = json.loads(raw)["data"]
    out = pd.DataFrame({"date": [d["date"] for d in data], "close": [float(d["close"]) for d in data]})
    out["date"] = pd.to_datetime(out["date"]).dt.strftime("%Y-%m-%d")
    return out.sort_values("date", kind="stable").reset_index(drop=True)


def history_frames(hist: dict[str, pd.DataFrame], column: str) -> pd.DataFrame:
    """Per-ticker yfinance history frames -> wide frame of one column, index `date` (date only, no time zone)."""
    cols = {}
    for ticker, frame in hist.items():
        s = frame[column].copy()
        idx = pd.DatetimeIndex(s.index)
        s.index = pd.Index((idx.tz_localize(None) if idx.tz is not None else idx).date, name="date")
        cols[ticker] = s
    out = pd.DataFrame(cols)
    out.index.name = "date"
    return out.sort_index()


def dividends_frame(hist: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """date, ticker, dividend for every nonzero dividend, in ticker order then date."""
    parts = []
    for ticker, frame in hist.items():
        d = frame.loc[frame["Dividends"] != 0, "Dividends"]
        idx = pd.DatetimeIndex(d.index)
        dates = (idx.tz_localize(None) if idx.tz is not None else idx).strftime("%Y-%m-%d")
        parts.append(pd.DataFrame({"date": dates, "ticker": ticker, "dividend": d.to_numpy(dtype=float)}))
    if not parts:
        return pd.DataFrame(columns=["date", "ticker", "dividend"])
    return pd.concat(parts, ignore_index=True)


def parse_xlk_holdings(raw: bytes) -> pd.DataFrame:
    """SSGA XLK workbook -> ticker, name, weight (decimal), shares, as_of_date."""
    ws = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    as_of = None
    for r in rows:
        if r and isinstance(r[0], str) and r[0].strip() == "Holdings:":
            as_of = datetime.strptime(str(r[1]).strip().removeprefix("As of").strip(), "%d-%b-%Y").date()
            break
    if as_of is None:
        raise ValueError("no 'Holdings:' line with 'As of <date>' in the workbook")
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Name")
    header = [str(c) if c is not None else "" for c in rows[head]]
    body = pd.DataFrame([list(r)[: len(header)] for r in rows[head + 1:]], columns=header)
    body = body[["Name", "Ticker", "Weight", "Shares Held"]]
    ticker = body["Ticker"].astype("string").str.strip()
    body = body[ticker.notna() & (ticker != "") & (ticker != "-")]
    return pd.DataFrame({
        "ticker": body["Ticker"].astype(str).str.strip(),
        "name": body["Name"].astype(str).str.strip(),
        "weight": body["Weight"].astype(float) / 100.0,  # workbook weight is in percent
        "shares": body["Shares Held"].astype(float),
        "as_of_date": as_of.isoformat(),
    }).reset_index(drop=True)


def check_xlk_weights(holdings: pd.DataFrame, raw: bytes) -> None:
    total = float(holdings["weight"].sum())
    if not XLK_WEIGHT_GATE[0] <= total <= XLK_WEIGHT_GATE[1]:
        ws = openpyxl.load_workbook(io.BytesIO(raw), read_only=True).worksheets[0]
        first = pd.DataFrame(list(ws.iter_rows(values_only=True, max_row=10)))
        raise GateError(f"XLK kept weights sum to {total:.10g}, outside {XLK_WEIGHT_GATE}\n{first.to_string()}")


# ---------------------------------------------------------------- manifest

def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def file_entry(path: Path, url: str, stage: str, pulled_utc: str) -> dict:
    rows = first = last = None
    if path.suffix == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False)
        rows = len(df)
        dates = pd.to_datetime(df.iloc[:, 0], format="mixed", errors="coerce").dropna()
        if df.columns[0].lower() in ("date", "observation_date") and len(dates):
            first, last = dates.min().strftime("%Y-%m-%d"), dates.max().strftime("%Y-%m-%d")
        elif "as_of_date" in df.columns and rows:
            first = last = df["as_of_date"].iloc[0]
    elif path.suffix == ".parquet":
        df = pd.read_parquet(path)
        rows = len(df)
        if rows:
            first, last = str(min(df.index)), str(max(df.index))
    elif path.suffix == ".xlsx":
        rows = openpyxl.load_workbook(path, read_only=True).worksheets[0].max_row
    return {"stage": stage, "url": url, "pulled_utc": pulled_utc, "rows": rows,
            "first_date": first, "last_date": last, "sha256": sha256(path)}


def write_manifest(stage: str, entries: dict[str, dict], pulled_utc: str, raw_root: Path = RAW) -> None:
    """Replace this stage's entries in MANIFEST.json, keeping the other stages'."""
    path = raw_root / "MANIFEST.json"
    manifest = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"stages": {}, "files": {}}
    manifest["files"] = {k: v for k, v in manifest["files"].items() if v["stage"] != stage}
    manifest["files"].update(entries)
    manifest["stages"][stage] = {
        "pulled_utc": pulled_utc,
        "versions": {"python": platform.python_version(), **{lib: version(lib) for lib in LIBS}},
    }
    manifest["files"] = dict(sorted(manifest["files"].items()))
    manifest["stages"] = dict(sorted(manifest["stages"].items()))
    path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n")


# ---------------------------------------------------------------- stages

def fetch(session, url: str) -> bytes:
    resp = session.get(url, timeout=60)
    resp.raise_for_status()
    if not resp.content:
        raise RuntimeError(f"empty response from {url}")
    return resp.content


def stage_cboe(session, pull_date: date, raw_root: Path = RAW) -> dict[str, dict]:
    out_dir = raw_root / "cboe"
    out_dir.mkdir(parents=True, exist_ok=True)
    pulled = datetime.now(UTC).isoformat()
    entries = {}
    for name in CBOE_CSV:
        url = CBOE_CSV_URL.format(name=name)
        path = out_dir / f"{name}.csv"
        path.write_bytes(fetch(session, url))  # saved unchanged
        entries[f"cboe/{name}.csv"] = file_entry(path, url, "cboe", pulled)
    for name in CBOE_JSON:
        url = CBOE_JSON_URL.format(name=name)
        path = out_dir / f"{name}.csv"
        cboe_json_to_frame(fetch(session, url)).to_csv(path, **CSV)
        entries[f"cboe/{name}.csv"] = file_entry(path, url, "cboe", pulled)
    write_manifest("cboe", entries, pulled, raw_root)
    return entries


def stage_fred(session, pull_date: date, raw_root: Path = RAW) -> dict[str, dict]:
    out_dir = raw_root / "fred"
    out_dir.mkdir(parents=True, exist_ok=True)
    pulled = datetime.now(UTC).isoformat()
    entries = {}
    for sid in FRED_IDS:
        url = FRED_URL.format(id=sid)
        path = out_dir / f"{sid}.csv"
        path.write_bytes(fetch(session, url))  # saved as served; missing values kept
        entries[f"fred/{sid}.csv"] = file_entry(path, url, "fred", pulled)
    write_manifest("fred", entries, pulled, raw_root)
    return entries


def history(ticker_factory, ticker: str, pull_date: date) -> pd.DataFrame:
    end = (pull_date + timedelta(days=1)).isoformat()  # yfinance end is exclusive
    hist = ticker_factory(ticker).history(start=PRICES_START, end=end, auto_adjust=False, actions=True)
    if hist is None or hist.empty:
        raise RuntimeError(f"yfinance returned no history for {ticker}")
    return hist


def stage_prices(ticker_factory, pull_date: date, cfg, raw_root: Path = RAW) -> dict[str, dict]:
    out_dir = raw_root / "prices"
    out_dir.mkdir(parents=True, exist_ok=True)
    pulled = datetime.now(UTC).isoformat()
    uni = cfg.universe
    captured = list(uni.etfs) + list(uni.xlk_holdings) + list(uni.singles)  # every captured ticker but ^SPX
    hist = {t: history(ticker_factory, t, pull_date) for t in captured}
    index_hist = {t: history(ticker_factory, t, pull_date) for t in INDEX_TICKERS}
    sector_hist = {t: history(ticker_factory, t, pull_date) for t in uni.sector_spdrs}
    url = f"yfinance Ticker(t).history(start={PRICES_START}, end={pull_date + timedelta(days=1)}, " \
          "auto_adjust=False, actions=True)"
    files = {
        "close.parquet": history_frames(hist, "Close"),
        "index.parquet": history_frames(index_hist, "Close"),
        "sectors.parquet": history_frames(sector_hist, "Adj Close"),
    }
    entries = {}
    for name, frame in files.items():
        frame.to_parquet(out_dir / name)
        entries[f"prices/{name}"] = file_entry(out_dir / name, url, "prices", pulled)
    dividends_frame(hist).to_csv(out_dir / "dividends.csv", **CSV)
    entries["prices/dividends.csv"] = file_entry(out_dir / "dividends.csv", url, "prices", pulled)
    write_manifest("prices", entries, pulled, raw_root)
    return entries


def stage_ssga(session, pull_date: date, raw_root: Path = RAW) -> dict[str, dict]:
    out_dir = raw_root / "ssga"
    out_dir.mkdir(parents=True, exist_ok=True)
    pulled = datetime.now(UTC).isoformat()
    raw = fetch(session, SSGA_URL)
    xlsx = out_dir / f"xlk_holdings_{pull_date.isoformat()}.xlsx"
    xlsx.write_bytes(raw)
    holdings = parse_xlk_holdings(raw)
    check_xlk_weights(holdings, raw)
    csv = out_dir / f"xlk_holdings_{pull_date.isoformat()}.csv"
    holdings.to_csv(csv, **CSV)
    entries = {f"ssga/{p.name}": file_entry(p, SSGA_URL, "ssga", pulled) for p in (xlsx, csv)}
    write_manifest("ssga", entries, pulled, raw_root)
    return entries


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--stage", choices=["cboe", "fred", "prices", "ssga", "all"], required=True)
    args = p.parse_args(argv)

    import requests
    import yfinance as yf

    cfg = load_config()
    pull_date = datetime.now(UTC).astimezone(NY).date()
    session = requests.Session()
    stages = ["cboe", "fred", "prices", "ssga"] if args.stage == "all" else [args.stage]
    for stage in stages:
        if stage == "prices":
            entries = stage_prices(yf.Ticker, pull_date, cfg)
        else:
            entries = {"cboe": stage_cboe, "fred": stage_fred, "ssga": stage_ssga}[stage](session, pull_date)
        for name, e in entries.items():
            print(f"{stage:6s} {name:36s} rows={e['rows']} {e['first_date']} .. {e['last_date']}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
